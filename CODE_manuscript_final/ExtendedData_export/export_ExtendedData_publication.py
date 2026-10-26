#!/usr/bin/env python3
"""Build publication Extended Data sheets from manuscript pipeline outputs.

Writes one workbook per sheet name under RESULT_manuscript_final/ExtendedData/,
matching the layout of ExtendedData.xlsx (caption row + header + data), and
rebuilds the combined ExtendedData.xlsx.

Inputs (defaults under openKoGES):
  DATA_manuscript_final/ExtendedData/
    ExtendedData1_continuous.csv, ExtendedData1_categorical.csv
    ExtendedData2_continuous.xlsx, ExtendedData2_categorical.xlsx
    ExtendedData_conditionally_independent_172.csv
    ExtendedData3.csv  (winner's-curse continuous analysis dump)
    ED1_phenotype_metadata.tsv

Usage:
  python3 export_ExtendedData_publication.py
"""

from __future__ import annotations

import json
import math
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils.dataframe import dataframe_to_rows
from scipy.optimize import minimize_scalar
from scipy.stats import norm

ROOT = Path("/home2/sjim/openKoGES")
OUT_DIR = ROOT / "RESULT_manuscript_final" / "ExtendedData"
INTERMEDIATE_DIR = ROOT / "DATA_manuscript_final" / "ExtendedData"
CODE_DIR = Path(__file__).resolve().parent
CAPTIONS_JSON = CODE_DIR / "sheet_captions.json"
TOC_TXT = CODE_DIR / "table_of_contents.txt"
ED1_META = INTERMEDIATE_DIR / "ED1_phenotype_metadata.tsv"

NA_TEXT = "Not analyzed"
NOT_APPLICABLE = "Not applicable"

# Continuous phenotype label aliases: metadata (publication) -> pipeline CSV
CONT_LABEL_ALIAS = {
    "LDL-cholesterol (LDL-C) *": "LDL-cholesterol (LDL-C)",
    "Phosphatae": "Phosphate",
    "Sytolic blood pressure (SBP)": "Systolic blood pressure (SBP)",
    "eGFR (CKD-EPI) ** ": "eGFR (CKD-EPI)",
}

ED2A_COLS = [
    "Phenotype",
    "Chromosome",
    "Base pair position (GRCh37)",
    "rsID",
    "Variant ID",
    "Nearest gene",
    "Effect allele",
    "Other allele",
    "KoGES beta",
    "KoGES standard error",
    "KoGES P value",
    "KoGES N",
    "Minimum P value within ±1 Mb in BBJ",
    "Minimum P value within ±1 Mb in TWB",
    "Minimum P value within ±1 Mb in TPMI",
    "EAS locus novelty",
    "Minimum P value within ±1 Mb in UKB",
    "Minimum P value within ±1 Mb in FinnGen",
    "Minimum P value within ±1 Mb in MVP",
    "EUR locus novelty",
    "GWAS Catalog novelty",
    "Integrated novelty status",
    "EAS panel meta analyzed",
    "EAS panel contributing cohorts",
    "EAS panel meta beta",
    "EAS panel meta standard error",
    "EAS panel meta P value",
    "EAS meta direction vs KoGES",
    "EAS meta nominal replication",
    "EAS meta Bonferroni replication",
    "EUR panel meta analyzed",
    "EUR panel contributing cohorts",
    "EUR panel meta beta",
    "EUR panel meta standard error",
    "EUR panel meta P value",
    "EUR meta direction vs KoGES",
    "EUR meta nominal replication",
    "EUR meta Bonferroni replication",
]

ED2B_COLS = [
    "Phenotype",
    "Chromosome",
    "Base pair position (GRCh37)",
    "rsID",
    "Variant ID",
    "Nearest gene",
    "Effect allele",
    "Other allele",
    "KoGES beta",
    "KoGES standard error",
    "KoGES P value",
    "KoGES N",
    "KoGES case",
    "KoGES control",
    "Minimum P value within ±1 Mb in BBJ",
    "Minimum P value within ±1 Mb in TPMI",
    "EAS locus novelty",
    "Minimum P value within ±1 Mb in UKB",
    "Minimum P value within ±1 Mb in FinnGen",
    "Minimum P value within ±1 Mb in MVP",
    "EUR locus novelty",
    "GWAS Catalog novelty",
    "Integrated novelty status",
    "EAS panel meta analyzed",
    "EAS panel contributing cohorts",
    "EAS panel meta beta",
    "EAS panel meta standard error",
    "EAS panel meta P value",
    "EAS meta direction vs KoGES",
    "EAS meta nominal replication",
    "EAS meta Bonferroni replication",
    "EUR panel meta analyzed",
    "EUR panel contributing cohorts",
    "EUR panel meta beta",
    "EUR panel meta standard error",
    "EUR panel meta P value",
    "EUR meta direction vs KoGES",
    "EUR meta nominal replication",
    "EUR meta Bonferroni replication",
]

ED3_COLS = [
    "Phenotype",
    "Lead variant (rsID)",
    "Variant ID",
    "Chromosome",
    "Base pair position (GRCh37)",
    "Variant type",
    "Nearest gene",
    "Effect allele",
    "Other allele",
    "Effect allele Frequency",
    "Minor Allele Frequency",
    "N",
    "Marginal beta",
    "Marginal SE",
    "Marginal P-value",
    "Conditional beta",
    "Conditional SE",
    "Conditional P-value",
    "n condition SNPs requested",
    "Condition SNPs retained by GCTA",
]

ED6A_COLS = [
    "Phenotype",
    "Variant ID",
    "rsID",
    "KoGES beta",
    "KoGES standard error",
    "KoGES P value",
    "KoGES N",
    "Novelty status",
    "Winner's curse corrected beta (β_WC)",
    "Shrinkage ratio (|β_WC/β_raw|)",
]

ED6B_COLS = [
    "Phenotype",
    "Variant ID",
    "rsID",
    "KoGES beta",
    "KoGES standard error",
    "KoGES P value",
    "KoGES N",
    "KoGES case",
    "KoGES control",
    "Novelty status",
    "Winner's curse corrected beta (β_WC)",
    "Shrinkage ratio (|β_WC/β_raw|)",
]


def load_captions() -> dict[str, str]:
    return json.loads(CAPTIONS_JSON.read_text(encoding="utf-8"))


def style_caption(cell) -> None:
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = Font(bold=True)


def write_sheet_workbook(
    path: Path,
    sheet_name: str,
    caption: str,
    header_rows: list[list],
    data: pd.DataFrame | None = None,
    extra_rows: list[list] | None = None,
) -> None:
    """Write a single-sheet workbook: caption, optional section headers, data."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append([caption])
    style_caption(ws["A1"])
    ws.row_dimensions[1].height = 80

    for row in header_rows:
        ws.append(row)

    if data is not None:
        for row in dataframe_to_rows(data, index=False, header=False):
            ws.append(list(row))

    if extra_rows:
        for row in extra_rows:
            ws.append(row)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def yes_no_repl(x) -> str:
    if pd.isna(x):
        return NA_TEXT
    if isinstance(x, str):
        s = x.strip().lower()
        if s in {"not analyzed", "na", ""}:
            return NA_TEXT
        if s in {"1", "true", "yes", "replicated"}:
            return "Replicated"
        if s in {"0", "false", "no", "not replicated"}:
            return "Not replicated"
        return x
    return "Replicated" if float(x) == 1 else "Not replicated"


def concordant_label(x) -> str:
    if pd.isna(x):
        return NA_TEXT
    if isinstance(x, str):
        s = x.strip().lower()
        if s in {"not analyzed", "na", ""}:
            return NA_TEXT
        if s in {"1", "true", "concordant"}:
            return "Concordant"
        if s in {"0", "false", "discordant"}:
            return "Discordant"
        return x
    return "Concordant" if float(x) == 1 else "Discordant"


def fill_na_text(df: pd.DataFrame, cols: list[str]) -> pd.DataFrame:
    out = df.copy()
    for c in cols:
        out[c] = out[c].where(out[c].notna(), NA_TEXT)
    return out


def build_ed2_summary(xlsx: Path, cols: list[str]) -> pd.DataFrame:
    df = pd.read_excel(xlsx, sheet_name="Total_summary_Novelty_Repl")
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise SystemExit(f"Missing columns in {xlsx.name}: {missing}")
    return df[cols].copy()


def build_panel_meta(
    xlsx: Path,
    ancestry: str,
) -> pd.DataFrame:
    """Format Replication_{EAS|EUR}_meta into Extended Data 4/5 publication columns."""
    sheet = f"Replication_{ancestry}_meta"
    src = pd.read_excel(xlsx, sheet_name=sheet)
    prefix = ancestry  # EAS or EUR

    analyzed = src["panel_analyzed"].fillna(False).astype(bool)
    # contributing_cohorts like "BBJ" -> 1; "BBJ+TWB" -> 2
    n_cohorts = src["contributing_cohorts"].map(
        lambda x: 0 if pd.isna(x) or str(x).strip() == "" else str(x).count("+") + 1
    )

    def pick(series, default=NA_TEXT):
        return np.where(analyzed, series, default)

    q = []
    i2 = []
    for ok, n, qv, iv in zip(
        analyzed,
        n_cohorts,
        src["external_cohort_Q_P"],
        src["external_cohort_I2"],
        strict=True,
    ):
        if not ok:
            q.append(NA_TEXT)
            i2.append(NA_TEXT)
        elif n < 2:
            q.append(NOT_APPLICABLE)
            i2.append(NOT_APPLICABLE)
        else:
            q.append(qv if pd.notna(qv) else NA_TEXT)
            i2.append(iv if pd.notna(iv) else NA_TEXT)

    out = pd.DataFrame(
        {
            "Phenotype": src["koges_phenotype"],
            "rsID": src["koges_lead_variant"],
            "Chromosome": src["koges_chr"],
            "Base pair position (GRCh37)": src["koges_bp"],
            "Contiributing cohorts": np.where(
                analyzed, src["contributing_cohorts"].fillna(NA_TEXT), NA_TEXT
            ),
            "KoGES beta": src["koges_beta_aligned_a1"],
            f"{prefix} panel meta beta": pick(src["external_panel_beta"]),
            f"{prefix} panel meta standard error": pick(src["external_panel_SE"]),
            f"{prefix} panel meta P value": pick(src["external_panel_P"]),
            "Direction concordant with KoGES": [
                concordant_label(v) if ok else NA_TEXT
                for ok, v in zip(analyzed, src["direction_concordant_with_KoGES"], strict=True)
            ],
            f"{prefix} meta nominal replication": [
                yes_no_repl(v) if ok else NA_TEXT
                for ok, v in zip(analyzed, src["replication_nominal"], strict=True)
            ],
            f"{prefix} meta Bonferroni replication": [
                yes_no_repl(v) if ok else NA_TEXT
                for ok, v in zip(analyzed, src["replication_Bonferroni"], strict=True)
            ],
            "Per-cohort direction pattern": np.where(
                analyzed,
                src["external_cohort_direction_pattern"].fillna(NA_TEXT),
                NA_TEXT,
            ),
            "Cochran Q P value": q,
            "I²": i2,
        }
    )
    # For non-analyzed rows, blank numeric KoGES beta still kept; meta fields already NA_TEXT
    for c in [
        f"{prefix} panel meta beta",
        f"{prefix} panel meta standard error",
        f"{prefix} panel meta P value",
    ]:
        out[c] = out[c].where(out[c].notna(), NA_TEXT)
    return out


def build_ed3(csv_path: Path) -> pd.DataFrame:
    src = pd.read_csv(csv_path)
    rename = {
        "EAF": "Effect allele Frequency",
        "MAF": "Minor Allele Frequency",
        "Marginal P": "Marginal P-value",
        "Conditional P": "Conditional P-value",
    }
    df = src.rename(columns=rename)
    missing = [c for c in ED3_COLS if c not in df.columns]
    if missing:
        raise SystemExit(f"ED3 missing columns: {missing}")
    return df[ED3_COLS].copy()


def wc_correct_one(beta_hat: float, p_val: float, z_cutoff: float) -> tuple[float, float]:
    """Return (wc_beta, shrinkage_abs_ratio). Same conditional-likelihood MLE as the R script."""
    if not np.isfinite(beta_hat) or beta_hat == 0:
        return (np.nan, np.nan)
    if not np.isfinite(p_val) or p_val < 0 or p_val > 1:
        return (np.nan, np.nan)
    # Match R: floor underflowed P=0 to .Machine$double.xmin
    p_val = max(float(p_val), np.finfo(float).tiny)
    z_mag = float(norm.ppf(1 - p_val / 2))
    if not np.isfinite(z_mag):
        # Extremely small P → ppf overflows; cap at a large finite Z
        z_mag = 40.0
    z_signed = math.copysign(z_mag, beta_hat)
    raw_z = abs(z_signed)
    if not np.isfinite(raw_z) or raw_z == 0:
        return (np.nan, np.nan)

    def nll(theta: float) -> float:
        p_selected = norm.cdf(-z_cutoff - theta) + norm.cdf(theta - z_cutoff)
        p_selected = max(p_selected, np.finfo(float).tiny)
        return -(norm.logpdf(raw_z, loc=theta, scale=1.0) - math.log(p_selected))

    upper = float(max(raw_z + 12, z_cutoff + 12))
    if not np.isfinite(upper) or upper <= 0:
        return (np.nan, np.nan)
    fit = minimize_scalar(nll, bounds=(0.0, upper), method="bounded")
    theta_wc = float(fit.x)
    shrink = theta_wc / raw_z if raw_z > 0 else np.nan
    return (beta_hat * shrink, shrink)


def build_ed6a_from_analysis_csv(csv_path: Path) -> pd.DataFrame:
    src = pd.read_csv(csv_path)
    out = pd.DataFrame(
        {
            "Phenotype": src["phenotype"],
            "Variant ID": src["variant_id"],
            "rsID": src["rsid"],
            "KoGES beta": src["koges_beta"],
            "KoGES standard error": src["koges_se"],
            "KoGES P value": src["koges_p_raw"].fillna(src["koges_p"]),
            "KoGES N": src["koges_n"],
            "Novelty status": src["integrated_novelty_status"],
            "Winner's curse corrected beta (β_WC)": src["wc_beta"],
            "Shrinkage ratio (|β_WC/β_raw|)": src["shrinkage_abs_ratio"],
        }
    )
    return out[ED6A_COLS]


def build_ed6b_from_cate_xlsx(xlsx: Path) -> pd.DataFrame:
    src = pd.read_excel(xlsx, sheet_name="Total_summary_Novelty_Repl")
    z_cut = norm.ppf(1 - 5e-8 / 2)
    betas = src["KoGES beta"].to_numpy(dtype=float)
    ps = src["KoGES P value"].to_numpy(dtype=float)
    wc_beta = []
    shrink = []
    for i, (b, p) in enumerate(zip(betas, ps, strict=True)):
        if (i + 1) % 100 == 0 or i + 1 == len(betas):
            print(f"  WC binary: {i + 1} / {len(betas)}")
        wb, sh = wc_correct_one(float(b), float(p), z_cut)
        wc_beta.append(wb)
        shrink.append(sh)
    out = pd.DataFrame(
        {
            "Phenotype": src["Phenotype"],
            "Variant ID": src["Variant ID"],
            "rsID": src["rsID"],
            "KoGES beta": src["KoGES beta"],
            "KoGES standard error": src["KoGES standard error"],
            "KoGES P value": src["KoGES P value"],
            "KoGES N": src["KoGES N"],
            "KoGES case": src["KoGES case"],
            "KoGES control": src["KoGES control"],
            "Novelty status": src["Integrated novelty status"],
            "Winner's curse corrected beta (β_WC)": wc_beta,
            "Shrinkage ratio (|β_WC/β_raw|)": shrink,
        }
    )
    return out[ED6B_COLS]


def build_ed1() -> list[list]:
    """Return section/header/data rows (without caption) for Extended Data1."""
    meta = pd.read_csv(ED1_META, sep="\t")
    cont = pd.read_csv(INTERMEDIATE_DIR / "ExtendedData1_continuous.csv")
    cate = pd.read_csv(INTERMEDIATE_DIR / "ExtendedData1_categorical.csv")

    cont_all = cont[cont["Group"] == "All"].set_index("Phenotype_label")
    cont_gwas = cont[cont["Group"] == "FAM_IDs"].set_index("Phenotype_label")
    cate_all = cate[cate["Group"] == "All"].set_index("Phenotype_label")
    cate_gwas = cate[cate["Group"] == "FAM_IDs"].set_index("Phenotype_label")

    rows: list[list] = []
    meta_c = meta[meta["section"] == "continuous"]
    rows.append([f"a. Continuous (Clinical Biomarkers) ({len(meta_c)})"])
    rows.append(
        [
            "Biomarker Category",
            "Phenotype",
            "Type",
            "Unit",
            "Total KoGES sample (N=211,562)",
            None,
            None,
            None,
            None,
            None,
            "GWAS sample (N=152,719)",
        ]
    )
    rows.append(
        [
            None,
            None,
            None,
            None,
            "Total(n)",
            "Mean",
            "Standard deviation",
            "Minimum",
            "Maximum",
            "NA(n)",
            "Total(n)",
            "Mean",
            "Standard deviation",
            "Minimum",
            "Maximum",
            "NA(n)",
        ]
    )

    prev_cat = None
    for _, m in meta_c.iterrows():
        pub_pheno = m["phenotype"]
        pipe_pheno = CONT_LABEL_ALIAS.get(pub_pheno, pub_pheno)
        if pipe_pheno not in cont_all.index:
            raise SystemExit(f"ED1 continuous phenotype not in CSV: {pipe_pheno}")
        a = cont_all.loc[pipe_pheno]
        g = cont_gwas.loc[pipe_pheno]
        cat = m["category"] if m["category"] != prev_cat else None
        prev_cat = m["category"]
        rows.append(
            [
                cat,
                pipe_pheno,
                "continuous",
                m["unit_or_definition"],
                a["Total(n)"],
                a["Mean"],
                a["SD"],
                a["Min"],
                a["Max"],
                a["NA(n)"],
                g["Total(n)"],
                g["Mean"],
                g["SD"],
                g["Min"],
                g["Max"],
                g["NA(n)"],
            ]
        )

    footnotes = [
        [
            "* Low-density lipoprotein cholesterol (LDL-C) was calculated using the "
            "Friedewald equation: LDL-C = Total cholesterol - HDL-C − (0.2×Triglycerides)"
        ],
        [
            "* eGFR was calculated using the CKD-EPI 2009 creatinine equation "
            "(Scr in mg/dL; 𝑘=0.7, 𝛼=−0.329 for women; 𝑘=0.9,𝛼=−0.411 for men)."
        ],
        [],
    ]

    meta_b = meta[meta["section"] == "binary"]
    rows_bin: list[list] = []
    rows_bin.append([f"b. Binary (Disease endpoints) ({len(meta_b)})"])
    rows_bin.append(
        [
            "Disease Category",
            "Phenotype",
            "Type",
            "Definition",
            "Total KoGES sample (N=211,562)",
            None,
            None,
            None,
            None,
            "GWAS sample (N=152,719)",
        ]
    )
    rows_bin.append(
        [
            None,
            None,
            None,
            None,
            "Total(n)",
            "Case(n)",
            "Control(n)",
            "Case-Control Ratio",
            "NA(n)",
            "Total(n)",
            "Case(n)",
            "Control(n)",
            "Case-Control Ratio",
            "NA(n)",
        ]
    )
    prev_cat = None
    for _, m in meta_b.iterrows():
        pheno = m["phenotype"]
        if pheno not in cate_all.index:
            raise SystemExit(f"ED1 binary phenotype not in CSV: {pheno}")
        a = cate_all.loc[pheno]
        g = cate_gwas.loc[pheno]
        cat = m["category"] if m["category"] != prev_cat else None
        prev_cat = m["category"]
        rows_bin.append(
            [
                cat,
                pheno,
                "binary",
                m["unit_or_definition"],
                a["Total(n)"],
                a["Case(n)"],
                a["Control(n)"],
                a["Case-Control Ratio"],
                a["NA(n)"],
                g["Total(n)"],
                g["Case(n)"],
                g["Control(n)"],
                g["Case-Control Ratio"],
                g["NA(n)"],
            ]
        )

    return rows + footnotes + rows_bin


def copy_sheet_from_gold(gold: Path, sheet_name: str) -> Workbook:
    """Preserve complex multi-header sheets (ED8) from the curated gold workbook."""
    src = load_workbook(gold, data_only=True)
    if sheet_name not in src.sheetnames:
        raise SystemExit(f"Gold workbook missing sheet: {sheet_name}")
    ws_src = src[sheet_name]
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    for row in ws_src.iter_rows(values_only=True):
        ws.append(list(row))
    src.close()
    return wb


def build_combined(sheet_files: dict[str, Path], captions: dict[str, str]) -> Path:
    wb = Workbook()
    # TOC
    ws = wb.active
    ws.title = "Table of Contents"
    if TOC_TXT.exists():
        for line in TOC_TXT.read_text(encoding="utf-8").splitlines():
            ws.append([line])
    else:
        for name, cap in captions.items():
            if name != "Table of Contents":
                ws.append([cap])

    for sheet_name, path in sheet_files.items():
        src = load_workbook(path, data_only=True)
        ws_src = src.active
        ws_new = wb.create_sheet(title=sheet_name[:31])
        for row in ws_src.iter_rows(values_only=True):
            ws_new.append(list(row))
        if sheet_name in captions:
            # ensure caption preserved
            pass
        src.close()

    out = OUT_DIR / "ExtendedData.xlsx"
    wb.save(out)
    return out


def main() -> None:
    captions = load_captions()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    sheet_files: dict[str, Path] = {}

    # --- Extended Data1 ---
    print("Building Extended Data1 ...")
    ed1_rows = build_ed1()
    path = OUT_DIR / "Extended Data1.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Extended Data1"
    ws.append([captions["Extended Data1"]])
    style_caption(ws["A1"])
    ws.row_dimensions[1].height = 90
    for row in ed1_rows:
        ws.append(row)
    wb.save(path)
    sheet_files["Extended Data1"] = path
    print("  ->", path)

    # --- Extended Data2A / 2B ---
    print("Building Extended Data2A/2B ...")
    for sheet, xlsx, cols in [
        ("Extended Data2A", INTERMEDIATE_DIR / "ExtendedData2_continuous.xlsx", ED2A_COLS),
        ("Extended Data2B", INTERMEDIATE_DIR / "ExtendedData2_categorical.xlsx", ED2B_COLS),
    ]:
        df = build_ed2_summary(xlsx, cols)
        path = OUT_DIR / f"{sheet}.xlsx"
        write_sheet_workbook(path, sheet, captions[sheet], [cols], df)
        sheet_files[sheet] = path
        print(f"  -> {path} ({len(df)} rows)")

    # --- Extended Data3 ---
    print("Building Extended Data3 ...")
    df = build_ed3(INTERMEDIATE_DIR / "ExtendedData_conditionally_independent_172.csv")
    path = OUT_DIR / "Extended Data3.xlsx"
    write_sheet_workbook(path, "Extended Data3", captions["Extended Data3"], [ED3_COLS], df)
    sheet_files["Extended Data3"] = path
    print(f"  -> {path} ({len(df)} rows)")

    # --- Extended Data4A/4B/5A/5B ---
    print("Building Extended Data4/5 panel meta ...")
    panel_specs = [
        ("Extended Data4A", INTERMEDIATE_DIR / "ExtendedData2_continuous.xlsx", "EAS"),
        ("Extended Data4B", INTERMEDIATE_DIR / "ExtendedData2_continuous.xlsx", "EUR"),
        ("Extended Data5A", INTERMEDIATE_DIR / "ExtendedData2_categorical.xlsx", "EAS"),
        ("Extended Data5B", INTERMEDIATE_DIR / "ExtendedData2_categorical.xlsx", "EUR"),
    ]
    for sheet, xlsx, ancestry in panel_specs:
        df = build_panel_meta(xlsx, ancestry)
        path = OUT_DIR / f"{sheet}.xlsx"
        write_sheet_workbook(
            path, sheet, captions[sheet], [list(df.columns)], df
        )
        sheet_files[sheet] = path
        print(f"  -> {path} ({len(df)} rows)")

    # --- Extended Data6A / 6B ---
    print("Building Extended Data6A ...")
    df = build_ed6a_from_analysis_csv(INTERMEDIATE_DIR / "ExtendedData3.csv")
    path = OUT_DIR / "Extended Data6A.xlsx"
    write_sheet_workbook(path, "Extended Data6A", captions["Extended Data6A"], [ED6A_COLS], df)
    sheet_files["Extended Data6A"] = path
    print(f"  -> {path} ({len(df)} rows)")

    print("Building Extended Data6B (WC on binary leads) ...")
    df = build_ed6b_from_cate_xlsx(INTERMEDIATE_DIR / "ExtendedData2_categorical.xlsx")
    path = OUT_DIR / "Extended Data6B.xlsx"
    write_sheet_workbook(path, "Extended Data6B", captions["Extended Data6B"], [ED6B_COLS], df)
    sheet_files["Extended Data6B"] = path
    print(f"  -> {path} ({len(df)} rows)")

    # --- Extended Data7A / 7B (from SuppleTable8_ExtendedData7.py outputs) ---
    print("Including Extended Data7A/7B if present ...")
    for sheet in ["Extended Data7A", "Extended Data7B"]:
        path = OUT_DIR / f"{sheet}.xlsx"
        if path.exists():
            sheet_files[sheet] = path
            print(f"  -> {path}")
        else:
            print(f"  [skip] missing {path.name} (run SuppleTable8_ExtendedData7.py first)")

    # --- Extended Data8A / 8B (preserve curated multi-header layout) ---
    gold = OUT_DIR / "ExtendedData.xlsx"
    print("Building Extended Data8A/8B from curated gold layout ...")
    for sheet in ["Extended Data8A", "Extended Data 8B"]:
        wb = copy_sheet_from_gold(gold, sheet)
        # filename mirrors sheet name
        path = OUT_DIR / f"{sheet}.xlsx"
        wb.save(path)
        sheet_files[sheet] = path
        print(f"  -> {path}")

    # Combined workbook (overwrite gold with regenerated sheets; keep ED8 content)
    print("Rebuilding combined ExtendedData.xlsx ...")
    out = build_combined(sheet_files, captions)
    print("  ->", out)
    print("Done.")


if __name__ == "__main__":
    main()
